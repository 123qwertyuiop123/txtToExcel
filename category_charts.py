"""月度和年度分类表、原生 Excel 图表的共用生成逻辑。

表格金额与图表绑定同一单元格范围；布局参数只影响位置和尺寸，
不影响分类口径、金额计算或“收入/支出分开”的规则。
"""

from dataclasses import dataclass
from decimal import Decimal
from math import ceil

from excel_utils import set_safe_text
from money_utils import decimal_to_number, round_one_decimal


PIE_CATEGORY_LIMIT = 8
EXPENSE_COLOR = "FCE4D6"
INCOME_COLOR = "E2F0D9"


@dataclass(frozen=True)
class CategoryLayout:
    """保留月份页与年度分类页各自的表格、图表尺寸。"""

    column: int
    title_size: int
    chart_width: int
    pie_height: int
    bar_min_height: int
    bar_max_height: int
    bar_extra_height: int


ANNUAL_LAYOUT = CategoryLayout(
    column=1, title_size=12, chart_width=13, pie_height=8,
    bar_min_height=9, bar_max_height=14, bar_extra_height=3,
)
MONTHLY_LAYOUT = CategoryLayout(
    column=5, title_size=11, chart_width=12, pie_height=7,
    bar_min_height=8, bar_max_height=12, bar_extra_height=2,
)


def chart_height(category_count: int, layout: CategoryLayout) -> float:
    """单位为厘米，与 openpyxl 的原生图表尺寸一致。"""
    if category_count <= PIE_CATEGORY_LIMIT:
        return layout.pie_height
    height = max(layout.bar_min_height, category_count * 0.45 + layout.bar_extra_height)
    return min(layout.bar_max_height, height)


def next_chart_row(sheet, start_row: int, category_count: int, layout: CategoryLayout) -> int:
    """按默认行高预留图表空间，避免大量分类时上下图表重叠。"""
    row_pixels = (sheet.sheet_format.defaultRowHeight or 15) * 96 / 72
    return start_row + ceil(chart_height(category_count, layout) * 96 / 2.54 / row_pixels) + 2


def add_category_section(
    sheet,
    values: dict[str, Decimal],
    *,
    start_row: int,
    table_title: str,
    chart_title: str,
    chart_anchor: str,
    color: str,
    layout: CategoryLayout,
    empty_message: str = "无可绘制数据",
) -> int:
    """写分类表和对应图表，返回总计行；无数据时只显示提示。"""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    column = layout.column
    amount_column = column + 1
    amount_letter = get_column_letter(amount_column)
    fill = PatternFill("solid", fgColor=color)
    title_cell = sheet.cell(start_row, column, table_title)
    title_cell.font = Font(bold=True, size=layout.title_size)
    title_cell.fill = fill
    sheet.merge_cells(start_row=start_row, start_column=column, end_row=start_row, end_column=column + 2)
    for offset, text in enumerate(("分类", "金额", "占比")):
        cell = sheet.cell(start_row + 1, column + offset, text)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    # 相同金额再按名称排序，确保重复转换产生稳定的表格顺序。
    sorted_values = sorted(values.items(), key=lambda item: (-item[1], item[0]))
    first_row = start_row + 2
    total_row = first_row + len(sorted_values)
    for row, (label, amount) in enumerate(sorted_values, start=first_row):
        set_safe_text(sheet.cell(row, column), label)
        sheet.cell(row, amount_column, decimal_to_number(round_one_decimal(amount)))
        sheet.cell(row, column + 2, f"=IFERROR({amount_letter}{row}/${amount_letter}${total_row},0)")
        sheet.cell(row, column + 2).number_format = "0.0%"

    sheet.cell(total_row, column, "总计" if sorted_values else "无数据")
    total = (
        f"=ROUND(SUM({amount_letter}{first_row}:{amount_letter}{total_row - 1}),1)"
        if sorted_values else 0
    )
    sheet.cell(total_row, amount_column, total)
    sheet.cell(total_row, column).font = Font(bold=True)
    sheet.cell(total_row, amount_column).font = Font(bold=True)

    if sorted_values:
        _add_chart(sheet, start_row, total_row, chart_title, chart_anchor, layout)
    else:
        sheet[chart_anchor] = empty_message
    return total_row


def _add_chart(sheet, start_row: int, total_row: int, title: str, anchor: str, layout: CategoryLayout) -> None:
    """共用图表标签设置：保留类别，只隐藏多余的系列名称“金额”。"""
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList

    first_row = start_row + 2
    category_count = total_row - first_row
    is_pie = category_count <= PIE_CATEGORY_LIMIT
    chart = PieChart() if is_pie else BarChart()
    if is_pie:
        chart.legend.position = "r"
    else:
        # 多分类时直接比较金额，占比仍可从旁边的公式表查看。
        chart.type = "bar"
        chart.style = 10
        chart.legend = None
        chart.x_axis.title = "金额"
        chart.x_axis.numFmt = "#,##0.0"
        chart.x_axis.scaling.min = 0
    chart.dataLabels = DataLabelList(
        showLegendKey=False,
        showVal=not is_pie,
        showCatName=True,
        showSerName=False,
        showPercent=is_pie,
        showBubbleSize=False,
        showLeaderLines=is_pie,
    )
    if not is_pie:
        chart.dataLabels.numFmt = "#,##0.0"
    chart.title = title
    chart.height = chart_height(category_count, layout)
    chart.width = layout.chart_width
    chart.add_data(
        Reference(sheet, min_col=layout.column + 1, min_row=start_row + 1, max_row=total_row - 1),
        titles_from_data=True,
    )
    chart.set_categories(Reference(sheet, min_col=layout.column, min_row=first_row, max_row=total_row - 1))
    sheet.add_chart(chart, anchor)
