"""按月份记录的普通消费 TXT 检查、解析和年度 Excel 生成。"""

import calendar
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from category_rules import categorize_expense
from excel_utils import save_workbook_safely, set_safe_text, validate_output_path
from models import DayRecord, MoneyItem
from money_utils import decimal_to_number, decimal_to_text, round_one_decimal


FILE_DATE_RE = re.compile(r"(?P<year>\d{4})\s*年\s*(?P<month>\d{1,2})\s*月")
DAY_RE = re.compile(r"^\s*(?P<day>\d{1,2})\s*[日号]?\s*[:：]\s*(?P<detail>.*)$")
ITEM_SPLIT_RE = re.compile(r"[,，、;；]")
ITEM_RE = re.compile(
    r"^\s*(?P<label>[^:：]+?)\s*[:：]\s*"
    r"(?P<amount>[+-]?\s*\d+(?:\.\d+)?)\s*$"
)
ITEM_WITHOUT_COLON_RE = re.compile(
    r"^\s*(?P<label>.*?\D)\s*(?P<amount>[+-]?\d+(?:\.\d+)?)\s*$"
)
RECOVERY_ITEM_RE = re.compile(
    r"(?P<label>[^\d+\-.,，、;；:：]+?)\s*[:：]?\s*"
    r"(?P<amount>[+-]?\d+(?:\.\d+)?)"
)


# ---------- 文本解析 ----------

def parse_money_items(detail: str) -> list[MoneyItem]:
    """解析明细，并恢复常见的漏写冒号、逗号情况。"""
    items: list[MoneyItem] = []

    def append_item(label: str, raw_amount: str) -> None:
        """完成单条金额转换，并保留“+”所表达的收入语义。"""
        amount_text = raw_amount.replace(" ", "")
        try:
            amount = Decimal(amount_text)
        except InvalidOperation as exc:
            raise ValueError(f"金额无法识别：{amount_text}") from exc
        items.append(MoneyItem(label.strip(), amount, amount_text.startswith("+")))

    segments = [segment.strip() for segment in ITEM_SPLIT_RE.split(detail) if segment.strip()]
    index = 0
    while index < len(segments):
        segment = segments[index]
        match = ITEM_RE.fullmatch(segment)
        if match is None and ":" not in segment and "：" not in segment:
            match = ITEM_WITHOUT_COLON_RE.fullmatch(segment)
        if match is not None:
            append_item(match.group("label"), match.group("amount"))
            index += 1
            continue

        # 兼容“午饭,14.7”：冒号被误写成逗号。
        if (
            not re.search(r"\d", segment)
            and index + 1 < len(segments)
            and re.fullmatch(r"[+-]?\d+(?:\.\d+)?", segments[index + 1])
        ):
            append_item(segment, segments[index + 1])
            index += 2
            continue

        # 兼容“晚饭:13,1”：小数点被误写成逗号。
        if (
            re.fullmatch(r"\d+", segment)
            and items
            and items[-1].amount == items[-1].amount.to_integral_value()
        ):
            previous = items[-1]
            items[-1] = MoneyItem(
                previous.label,
                Decimal(f"{previous.amount}.{segment}"),
                previous.is_income,
            )
            index += 1
            continue

        # 最后尝试从漏写分隔符的连续文本中恢复多条明细。
        recovery_matches = list(RECOVERY_ITEM_RE.finditer(segment))
        unparsed = RECOVERY_ITEM_RE.sub("", segment).strip()
        if not recovery_matches or unparsed:
            raise ValueError(f"无法识别明细项目：{segment.strip()}")
        for recovery_match in recovery_matches:
            append_item(recovery_match.group("label"), recovery_match.group("amount"))
        index += 1
    return items


def parse_txt(path: Path) -> dict[int, DayRecord]:
    """读取单个月份 TXT，并按日期返回记录。"""
    records: dict[int, DayRecord] = {}
    text = path.read_text(encoding="utf-8-sig")
    for line_number, raw_line in enumerate(text.splitlines(), start=1):
        if not raw_line.strip():
            continue
        match = DAY_RE.match(raw_line)
        if not match:
            raise ValueError(f"{path.name} 第 {line_number} 行格式无法识别：{raw_line}")
        day = int(match.group("day"))
        if not 1 <= day <= 31:
            raise ValueError(f"文件：{path.name}\n第 {line_number} 行\n日期必须为 1–31 日：{day}日")
        detail = match.group("detail").strip()
        if day in records:
            raise ValueError(f"{path.name} 中第 {day} 日重复出现")
        try:
            items = parse_money_items(detail)
        except ValueError as exc:
            raise ValueError(f"文件：{path.name}\n第 {line_number} 行（{day}日）\n{exc}") from exc
        records[day] = DayRecord(day, detail, items)
    return records


def infer_year_month(path: Path, year: int | None, month: int | None) -> tuple[int, int]:
    """优先使用手动年月，否则从文件名识别。"""
    match = FILE_DATE_RE.search(path.stem)
    inferred_year = int(match.group("year")) if match else None
    inferred_month = int(match.group("month")) if match else None
    final_year = year if year is not None else inferred_year
    final_month = month if month is not None else inferred_month
    if final_year is None or final_month is None:
        raise ValueError(f"无法从文件名“{path.name}”识别年月，请使用 --year 和 --month 指定")
    if not 1 <= final_month <= 12:
        raise ValueError(f"月份必须为 1 到 12，实际为 {final_month}")
    if not 1 <= final_year <= 9999:
        raise ValueError(f"年份必须为 1 到 9999，实际为 {final_year}")
    return final_year, final_month


def find_txt_files(input_path: Path) -> list[Path]:
    """把单个文件或目录统一展开为 TXT 文件列表。"""
    if input_path.is_file():
        if input_path.suffix.lower() != ".txt":
            raise ValueError(f"请选择 .txt 文件：{input_path.name}")
        return [input_path]
    if input_path.is_dir():
        files = sorted(input_path.glob("*.txt"))
        if files:
            return files
        raise FileNotFoundError(f"目录中没有 TXT 文件：{input_path}")
    raise FileNotFoundError(f"输入路径不存在：{input_path}")


# ---------- 严格检查 ----------

def validate_txt_standard(path: Path, year: int, month: int) -> list[str]:
    """严格检查单个文件；检查阶段不使用自动容错，以便用户修正源文件。"""
    errors: list[str] = []
    try:
        text = path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError) as exc:
        return [f"文件：{path.name}\n无法读取 UTF-8 文本：{exc}"]

    max_day = calendar.monthrange(year, month)[1]
    seen_days: dict[int, int] = {}
    for line_number, raw_line in enumerate(text.splitlines(), start=1):
        if not raw_line.strip():
            continue
        day_match = DAY_RE.match(raw_line)
        if not day_match:
            errors.append(
                f"文件：{path.name}，第 {line_number} 行\n"
                f"日期格式不正确：{raw_line}\n应写为：1日:早饭:8"
            )
            continue
        day = int(day_match.group("day"))
        if not 1 <= day <= max_day:
            errors.append(
                f"文件：{path.name}，第 {line_number} 行（{day}日）\n"
                f"日期超出 {year}年{month}月的有效范围（1–{max_day}日）"
            )
        if day in seen_days:
            errors.append(
                f"文件：{path.name}，第 {line_number} 行（{day}日）\n"
                f"日期重复，首次出现在第 {seen_days[day]} 行"
            )
        else:
            seen_days[day] = line_number

        detail = day_match.group("detail").strip()
        if not detail:
            continue
        for segment in ITEM_SPLIT_RE.split(detail):
            item = segment.strip()
            if not item:
                errors.append(
                    f"文件：{path.name}，第 {line_number} 行（{day}日）\n"
                    "存在多余或连续的分隔符，请删除空白项目"
                )
            elif ITEM_RE.fullmatch(item) is None:
                errors.append(
                    f"文件：{path.name}，第 {line_number} 行（{day}日）\n"
                    f"明细格式不标准：{item}\n应写为：项目:金额，例如 晚饭:10"
                )
    return errors


def validate_many(input_paths: Iterable[Path], year: int | None, month: int | None) -> list[str]:
    """检查多个文件是否可以安全合并到同一年度工作簿。"""
    errors: list[str] = []
    files: list[Path] = []
    seen_paths: set[Path] = set()
    for input_path in input_paths:
        try:
            discovered = find_txt_files(input_path)
        except (OSError, ValueError) as exc:
            errors.append(str(exc))
            continue
        for file in discovered:
            resolved = file.resolve()
            if resolved not in seen_paths:
                files.append(file)
                seen_paths.add(resolved)
    if not files:
        return errors or ["没有选择 TXT 文件"]
    if len(files) > 1 and month is not None:
        errors.append("选择多个文件时，请通过文件名标明各自月份，不要手动指定统一月份")

    workbook_year: int | None = None
    year_source: Path | None = None
    month_sources: dict[int, Path] = {}
    file_dates: list[tuple[Path, int, int]] = []
    for file in files:
        try:
            file_year, file_month = infer_year_month(
                file,
                year if len(files) == 1 else None,
                month if len(files) == 1 else None,
            )
        except ValueError as exc:
            errors.append(str(exc))
            continue
        if workbook_year is not None and file_year != workbook_year:
            errors.append(
                "所选文件不属于同一年：\n"
                f"{year_source.name if year_source else '未知'}（{workbook_year}年）\n"
                f"{file.name}（{file_year}年）"
            )
        elif workbook_year is None:
            workbook_year = file_year
            year_source = file
        if file_month in month_sources:
            errors.append(
                f"月份重复：{file_month}月\n"
                f"{month_sources[file_month].name}\n{file.name}"
            )
        else:
            month_sources[file_month] = file
        file_dates.append((file, file_year, file_month))

    for file, file_year, file_month in file_dates:
        errors.extend(validate_txt_standard(file, file_year, file_month))
    return errors


def expense_detail(record: DayRecord) -> str:
    """生成只包含支出的明细文本；收入统一移动到月份底部。"""
    return ",".join(
        f"{item.label}:{decimal_to_text(item.amount)}"
        for item in record.items
        if not item.is_income
    )


def _aggregate_categories(
    month_records: dict[int, dict[int, DayRecord]],
) -> tuple[dict[str, Decimal], dict[str, Decimal]]:
    """分别汇总全年支出原因和收入来源。"""
    expenses: dict[str, Decimal] = {}
    incomes: dict[str, Decimal] = {}
    for records in month_records.values():
        for record in records.values():
            for item in record.items:
                # 收入来源保持原名；支出只对命中明确规则的原因进行合并。
                label = item.label if item.is_income else categorize_expense(item.label)
                target = incomes if item.is_income else expenses
                target[label] = target.get(label, Decimal("0")) + abs(item.amount)
    return expenses, incomes


def _create_category_sheet(workbook, month_records: dict[int, dict[int, DayRecord]], year: int) -> None:
    """年度分类页只定义布局，表格与图表生成交给共用模块。"""
    from category_charts import (
        ANNUAL_LAYOUT, EXPENSE_COLOR, INCOME_COLOR, PIE_CATEGORY_LIMIT, add_category_section,
    )

    expenses, incomes = _aggregate_categories(month_records)
    sheet = workbook.create_sheet("分类统计")
    sheet.sheet_view.showGridLines = True
    sheet.freeze_panes = "A3"
    total_row = add_category_section(
        sheet, expenses, start_row=1, table_title=f"{year}年支出原因占比",
        chart_title=f"{year}年支出原因占比", chart_anchor="E2",
        color=EXPENSE_COLOR, layout=ANNUAL_LAYOUT, empty_message="本年度无可绘制数据",
    )
    income_row = max(31 if len(expenses) > PIE_CATEGORY_LIMIT else 21, total_row + 3)
    add_category_section(
        sheet, incomes, start_row=income_row, table_title=f"{year}年收入来源占比",
        chart_title=f"{year}年收入来源占比", chart_anchor=f"E{income_row + 1}",
        color=INCOME_COLOR, layout=ANNUAL_LAYOUT, empty_message="本年度无可绘制数据",
    )
    for column, width in {"A": 24, "B": 16, "C": 14}.items():
        sheet.column_dimensions[column].width = width


def _add_monthly_category_charts(sheet, records: dict[int, DayRecord], year: int, month: int) -> None:
    """只在月份页右侧增加统计内容，保持 A:C 明细与原有样式不变。"""
    from category_charts import (
        EXPENSE_COLOR, INCOME_COLOR, MONTHLY_LAYOUT, add_category_section, next_chart_row,
    )

    expenses, incomes = _aggregate_categories({month: records})
    total_row = add_category_section(
        sheet, expenses, start_row=1, table_title=f"{month}月支出占比",
        chart_title=f"{year}年{month}月支出占比", chart_anchor="I2",
        color=EXPENSE_COLOR, layout=MONTHLY_LAYOUT,
    )
    income_row = max(20, total_row + 3, next_chart_row(sheet, 2, len(expenses), MONTHLY_LAYOUT))
    add_category_section(
        sheet, incomes, start_row=income_row, table_title=f"{month}月收入占比",
        chart_title=f"{year}年{month}月收入占比", chart_anchor=f"I{income_row + 1}",
        color=INCOME_COLOR, layout=MONTHLY_LAYOUT,
    )
    for column, width in {"E": 18, "F": 14, "G": 12}.items():
        sheet.column_dimensions[column].width = width


# ---------- Excel 生成 ----------

def _validate_month_records(records: dict[int, DayRecord], year: int, month: int) -> None:
    """生成前验证完整日期；不能让循环月份天数时静默丢弃非法记录。"""
    if not 1 <= year <= 9999 or not 1 <= month <= 12:
        raise ValueError(f"年月无效：{year}年{month}月")
    max_day = calendar.monthrange(year, month)[1]
    for day, record in records.items():
        if not 1 <= day <= max_day:
            raise ValueError(f"日期超出 {year}年{month}月的有效范围（1–{max_day}日）：{day}日")
        if record.day != day:
            raise ValueError(f"日期键与记录不一致：{day}日 / {record.day}日")


def create_workbook(
    month_records: dict[int, dict[int, DayRecord]],
    year: int,
    output_path: Path,
) -> None:
    """生成包含 12 个月份和年度总计的普通消费工作簿。"""
    _validate_month_records({}, year, 1)
    for month_number, records in month_records.items():
        _validate_month_records(records, year, month_number)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请先执行：pip install -r requirements.txt") from exc

    workbook = Workbook()
    workbook.remove(workbook.active)
    total_rows: dict[int, tuple[int, int]] = {}
    for month_number in range(1, 13):
        sheet = workbook.create_sheet(f"{month_number}月")
        sheet.sheet_view.showGridLines = True
        records = month_records.get(month_number, {})
        days_in_month = calendar.monthrange(year, month_number)[1]
        income_parts: list[str] = []
        income_values: list[str] = []
        for day in range(1, days_in_month + 1):
            record = records.get(day, DayRecord(day, ""))
            sheet.cell(day, 1, f"{day}日:")
            # 外部文字统一强制写成字符串，避免以公式前缀开头时被 Excel 执行。
            set_safe_text(sheet.cell(day, 2), expense_detail(record) or None)
            if record.expense:
                sheet.cell(day, 3, decimal_to_number(round_one_decimal(record.expense)))
            for item in record.items:
                if item.is_income:
                    value = decimal_to_number(abs(item.amount))
                    income_parts.append(f"{item.label}:+{value}")
                    income_values.append(str(value))

        expense_row = days_in_month + 1
        income_row = days_in_month + 2
        sheet.cell(expense_row, 2, "总计")
        sheet.cell(expense_row, 3, f"=ROUND(SUM(C1:C{days_in_month}),1)")
        set_safe_text(
            sheet.cell(income_row, 2),
            "，".join(income_parts) if income_parts else "收入总计",
        )
        sheet.cell(income_row, 3, "=" + "+".join(income_values) if income_values else "=0")
        total_rows[month_number] = (expense_row, income_row)
        sheet.column_dimensions["A"].width = 10
        sheet.column_dimensions["B"].width = 46
        sheet.column_dimensions["C"].width = 16
        for row in sheet.iter_rows(min_row=1, max_row=income_row, min_col=1, max_col=3):
            for cell in row:
                cell.alignment = Alignment(vertical="center")
        sheet.cell(expense_row, 3).font = Font(bold=True)
        _add_monthly_category_charts(sheet, records, year, month_number)

    summary = workbook.create_sheet("总计")
    summary.append(["月份", "支出", "收入", "结余（收入-支出）"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for month_number in range(1, 13):
        row = month_number + 1
        expense_row, income_row = total_rows[month_number]
        summary.cell(row, 1, f"{month_number}月")
        summary.cell(row, 2, f"='{month_number}月'!C{expense_row}")
        summary.cell(row, 3, f"='{month_number}月'!C{income_row}")
        summary.cell(row, 4, f"=ROUND(C{row}-B{row},1)")
    final_row = 14
    summary.cell(final_row, 1, f"{year}年总计")
    summary.cell(final_row, 2, "=ROUND(SUM(B2:B13),1)")
    summary.cell(final_row, 3, "=ROUND(SUM(C2:C13),1)")
    summary.cell(final_row, 4, f"=ROUND(C{final_row}-B{final_row},1)")
    for cell in summary[final_row]:
        cell.font = Font(bold=True)
    for column, width in {"A": 14, "B": 18, "C": 18, "D": 24}.items():
        summary.column_dimensions[column].width = width
    # 年度分类另设统计页；月份页仅在右侧添加图表，总计页仍保持原布局。
    _create_category_sheet(workbook, month_records, year)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    # 先完整写入同目录临时文件，再替换目标，避免中断时损坏旧工作簿。
    save_workbook_safely(workbook, output_path)


def convert(input_path: Path, output_path: Path | None, year: int | None, month: int | None) -> Path:
    """单文件转换入口，内部复用多文件转换流程。"""
    return convert_many([input_path], output_path, year, month)


def convert_many(
    input_paths: Iterable[Path],
    output_path: Path | None,
    year: int | None,
    month: int | None,
) -> Path:
    """将一个或多个月份 TXT 转为同一个年度工作簿。"""
    files: list[Path] = []
    seen: set[Path] = set()
    for input_path in input_paths:
        for file in find_txt_files(input_path):
            resolved = file.resolve()
            if resolved not in seen:
                files.append(file)
                seen.add(resolved)
    if not files:
        raise FileNotFoundError("没有选择 TXT 文件")
    if len(files) > 1 and month is not None:
        raise ValueError("选择多个文件时，月份必须从各文件名中识别")

    month_records: dict[int, dict[int, DayRecord]] = {}
    month_sources: dict[int, Path] = {}
    workbook_year: int | None = None
    first_year_file: Path | None = None
    for path in files:
        file_year, file_month = infer_year_month(
            path,
            year if len(files) == 1 else None,
            month if len(files) == 1 else None,
        )
        if workbook_year is not None and file_year != workbook_year:
            raise ValueError(
                "一次只能生成一个年份的工作簿，请按年份分开转换。\n"
                f"文件：{first_year_file.name if first_year_file else '未知'}（{workbook_year}年）\n"
                f"文件：{path.name}（{file_year}年）"
            )
        if file_month in month_records:
            raise ValueError(
                f"存在多个 {file_month} 月文件，无法确定使用哪一个：\n"
                f"{month_sources[file_month].name}\n{path.name}"
            )
        workbook_year = file_year
        if first_year_file is None:
            first_year_file = path
        records = parse_txt(path)
        try:
            _validate_month_records(records, file_year, file_month)
        except ValueError as exc:
            raise ValueError(f"文件：{path.name}\n{exc}") from exc
        month_records[file_month] = records
        month_sources[file_month] = path

    assert workbook_year is not None
    target = validate_output_path(output_path or files[0].parent / f"{workbook_year}年消费统计.xlsx", files)
    create_workbook(month_records, workbook_year, target)
    return target
