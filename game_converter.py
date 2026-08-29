"""年度游戏消费 TXT 的检查、解析和 Excel 生成。"""

import re
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from excel_utils import save_workbook_safely, set_safe_text
from models import GameRecord
from money_utils import decimal_to_number, round_one_decimal


# 文件名决定所属年度；行内年份可省略，也允许截图中的“年:”写法。
GAME_FILE_RE = re.compile(r"^(?P<year>\d{4})\s*年\s*游戏$")
GAME_LINE_RE = re.compile(
    r"^\s*(?:(?P<year>\d{4})\s*年\s*[:：]?\s*)?"
    r"(?P<month>\d{1,2})\s*月\s*(?P<day>\d{1,2})\s*日\s*[:：]\s*"
    r"(?P<game>[^:：]+?)\s*[:：]\s*(?P<amount>\d+(?:\.\d+)?)\s*$"
)


# ---------- 文件名和文本解析 ----------

def infer_game_year(path: Path) -> int:
    """从“2025年游戏.txt”这类标准文件名中提取年份。"""
    match = GAME_FILE_RE.fullmatch(path.stem)
    if match is None:
        raise ValueError(f"文件名不符合标准：{path.name}\n应命名为：2025年游戏.txt")
    return int(match.group("year"))


def _parse_game_line(path: Path, line_number: int, raw_line: str, file_year: int) -> GameRecord:
    """解析一行，并生成包含文件名和行号的错误信息。"""
    match = GAME_LINE_RE.fullmatch(raw_line)
    if match is None:
        raise ValueError(
            f"文件：{path.name}，第 {line_number} 行\n"
            f"内容格式不正确：{raw_line}\n"
            "应写为：1月29日:原神:68 或 2025年1月29日:原神:68"
        )
    line_year = int(match.group("year")) if match.group("year") else file_year
    if line_year != file_year:
        raise ValueError(
            f"文件：{path.name}，第 {line_number} 行\n"
            f"行内年份为 {line_year} 年，与文件名中的 {file_year} 年不一致"
        )

    month = int(match.group("month"))
    day = int(match.group("day"))
    # date 会统一检查闰年、大小月和日期范围。
    try:
        purchase_date = date(file_year, month, day)
    except ValueError as exc:
        raise ValueError(
            f"文件：{path.name}，第 {line_number} 行\n"
            f"日期无效：{file_year}年{month}月{day}日"
        ) from exc
    try:
        amount = Decimal(match.group("amount"))
    except InvalidOperation as exc:
        raise ValueError(f"文件：{path.name}，第 {line_number} 行\n金额无法识别") from exc
    return GameRecord(purchase_date, match.group("game").strip(), amount)


def parse_game_txt(path: Path, year: int | None = None) -> list[GameRecord]:
    """读取一个年度游戏 TXT；空行跳过，非空行必须全部可识别。"""
    file_year = year if year is not None else infer_game_year(path)
    records: list[GameRecord] = []
    text = path.read_text(encoding="utf-8-sig")
    for line_number, raw_line in enumerate(text.splitlines(), start=1):
        if raw_line.strip():
            records.append(_parse_game_line(path, line_number, raw_line, file_year))
    if not records:
        raise ValueError(f"文件：{path.name}\n文件中没有游戏消费记录")
    return records


# ---------- 转换前检查 ----------

def validate_game_txt_files(paths: Iterable[Path]) -> list[str]:
    """一次收集全部游戏 TXT 问题，方便用户集中修改。"""
    errors: list[str] = []
    seen_years: dict[int, Path] = {}
    selected = list(paths)
    if not selected:
        return ["没有选择游戏消费 TXT 文件"]
    for path in selected:
        if path.suffix.lower() != ".txt":
            errors.append(f"文件格式不正确：{path.name}\n请选择 .txt 文件")
            continue
        if not path.is_file():
            errors.append(f"文件不存在：{path}")
            continue
        try:
            file_year = infer_game_year(path)
        except ValueError as exc:
            errors.append(str(exc))
            continue
        # 一个年度只允许一个来源文件，防止同一年数据被重复计入。
        if file_year in seen_years:
            errors.append(f"年份重复：{file_year}年\n{seen_years[file_year].name}\n{path.name}")
            continue
        seen_years[file_year] = path
        try:
            text = path.read_text(encoding="utf-8-sig")
        except (OSError, UnicodeError) as exc:
            errors.append(f"文件：{path.name}\n无法读取 UTF-8 文本：{exc}")
            continue

        record_count = 0
        for line_number, raw_line in enumerate(text.splitlines(), start=1):
            if not raw_line.strip():
                continue
            record_count += 1
            try:
                _parse_game_line(path, line_number, raw_line, file_year)
            except ValueError as exc:
                errors.append(str(exc))
        if record_count == 0:
            errors.append(f"文件：{path.name}\n文件中没有游戏消费记录")
    return errors


# ---------- Excel 生成 ----------

def create_game_workbook(records_by_year: dict[int, list[GameRecord]], output_path: Path) -> None:
    """生成总览，以及每年的逐笔、月度和游戏分类数据。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请先执行：pip install -r requirements.txt") from exc

    workbook = Workbook()
    summary = workbook.active
    summary.title = "游戏统计"
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    summary.append(["年份", "游戏消费总计"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # 年份排序保证工作表和总表始终按时间顺序排列，与选择文件的顺序无关。
    for summary_row, year in enumerate(sorted(records_by_year), start=2):
        records = sorted(records_by_year[year], key=lambda item: (item.purchase_date, item.game))
        sheet = workbook.create_sheet(f"{year}年")
        sheet.freeze_panes = "A2"
        # D、G 列分隔逐笔明细、月度汇总和游戏分类汇总。
        sheet.append(["日期", "游戏", "金额", None, "月份", "月消费", None, "游戏", "游戏消费"])
        for column in (1, 2, 3, 5, 6, 8, 9):
            cell = sheet.cell(1, column)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        # Decimal 避免浮点数累计产生金额误差。
        monthly_totals = {month: Decimal("0") for month in range(1, 13)}
        game_totals: dict[str, Decimal] = {}
        # 明细写入和两类汇总在一次遍历中完成，避免重复扫描全部记录。
        for row, record in enumerate(records, start=2):
            sheet.cell(row, 1, record.purchase_date)
            sheet.cell(row, 1).number_format = 'm"月"d"日"'
            set_safe_text(sheet.cell(row, 2), record.game)
            sheet.cell(row, 3, decimal_to_number(round_one_decimal(record.amount)))
            monthly_totals[record.purchase_date.month] += record.amount
            game_totals[record.game] = game_totals.get(record.game, Decimal("0")) + record.amount

        for month in range(1, 13):
            row = month + 1
            sheet.cell(row, 5, f"{month}月")
            sheet.cell(row, 6, decimal_to_number(round_one_decimal(monthly_totals[month])))
        total = sum((record.amount for record in records), Decimal("0"))
        sheet.cell(14, 5, f"{year}年总计")
        sheet.cell(14, 6, decimal_to_number(round_one_decimal(total)))
        sheet.cell(14, 5).font = Font(bold=True)
        sheet.cell(14, 6).font = Font(bold=True)

        for row, (game, game_total) in enumerate(sorted(game_totals.items()), start=2):
            set_safe_text(sheet.cell(row, 8), game)
            sheet.cell(row, 9, decimal_to_number(round_one_decimal(game_total)))
        for column, width in {
            "A": 14, "B": 20, "C": 14, "D": 3, "E": 12,
            "F": 14, "G": 3, "H": 20, "I": 14,
        }.items():
            sheet.column_dimensions[column].width = width
        sheet.auto_filter.ref = f"A1:C{len(records) + 1}"
        summary.cell(summary_row, 1, f"{year}年")
        summary.cell(summary_row, 2, decimal_to_number(round_one_decimal(total)))

    final_row = len(records_by_year) + 2
    summary.cell(final_row, 1, "全部年份总计")
    summary.cell(final_row, 2, f"=ROUND(SUM(B2:B{final_row - 1}),1)")
    summary.cell(final_row, 1).font = Font(bold=True)
    summary.cell(final_row, 2).font = Font(bold=True)
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 20
    summary.freeze_panes = "A2"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    save_workbook_safely(workbook, output_path)


def convert_game_txt_files(paths: Iterable[Path], output_path: Path) -> Path:
    """检查并转换多个年度游戏 TXT；任何文件失败时都不生成结果。"""
    selected = list(paths)
    errors = validate_game_txt_files(selected)
    if errors:
        raise ValueError("游戏消费 TXT 检查未通过：\n\n" + "\n\n".join(errors))
    records_by_year = {infer_game_year(path): parse_game_txt(path) for path in selected}
    create_game_workbook(records_by_year, output_path)
    return output_path
