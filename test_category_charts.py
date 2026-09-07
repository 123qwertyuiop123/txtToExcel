import unittest
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from category_rules import categorize_expense
from expense_converter import create_workbook
from models import DayRecord, MoneyItem


class CategoryChartTests(unittest.TestCase):
    def test_only_clear_expense_labels_are_merged(self):
        for label in ("早饭", "晚饭", "KFC", "饮料", "烧烤"):
            self.assertEqual(categorize_expense(label), "餐饮")
        self.assertEqual(categorize_expense("火车票"), "交通")
        self.assertEqual(categorize_expense("原神"), "游戏消费")
        self.assertEqual(categorize_expense("挂号"), "医疗")
        self.assertEqual(categorize_expense("宽带"), "网络通信")
        # 含义不明确的原因必须原样保留。
        for label in ("购物", "生鲜", "网吧", "裁缝", "妈妈"):
            self.assertEqual(categorize_expense(label), label)

    def test_expense_and_income_have_separate_tables_and_charts(self):
        records = {
            1: {
                1: DayRecord(1, "", [
                    MoneyItem("早餐", Decimal("10"), False),
                    MoneyItem("工资", Decimal("100"), True),
                ]),
                2: DayRecord(2, "", [
                    MoneyItem("早餐", Decimal("5"), False),
                    MoneyItem("交通", Decimal("20"), False),
                    MoneyItem("奖金", Decimal("50"), True),
                ]),
            }
        }
        with TemporaryDirectory() as directory:
            output = Path(directory) / "result.xlsx"
            create_workbook(records, 2025, output)
            workbook = load_workbook(output, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames[-1], "分类统计")
                sheet = workbook["分类统计"]
                self.assertEqual(len(sheet._charts), 2)
                self.assertEqual(sheet["A1"].value, "2025年支出原因占比")
                self.assertEqual(sheet["A3"].value, "交通")
                self.assertEqual(sheet["B3"].value, 20)
                self.assertEqual(sheet["A4"].value, "餐饮")
                self.assertEqual(sheet["B4"].value, 15)
                self.assertEqual(sheet["C3"].value, "=IFERROR(B3/$B$5,0)")
                self.assertEqual(sheet["B5"].value, "=ROUND(SUM(B3:B4),1)")
                self.assertEqual(sheet["A21"].value, "2025年收入来源占比")
                self.assertEqual(sheet["A23"].value, "工资")
                self.assertEqual(sheet["B23"].value, 100)

                monthly = workbook["1月"]
                self.assertEqual(len(monthly._charts), 2)
                self.assertEqual(monthly["E1"].value, "1月支出占比")
                self.assertEqual(monthly["E3"].value, "交通")
                self.assertEqual(monthly["F3"].value, 20)
                self.assertEqual(monthly["E20"].value, "1月收入占比")
                self.assertEqual(monthly["E22"].value, "工资")
                self.assertEqual(monthly["F22"].value, 100)
            finally:
                workbook.close()

    def test_no_income_does_not_create_blank_income_chart(self):
        records = {
            1: {1: DayRecord(1, "", [MoneyItem("早餐", Decimal("10"), False)])}
        }
        with TemporaryDirectory() as directory:
            output = Path(directory) / "result.xlsx"
            create_workbook(records, 2025, output)
            workbook = load_workbook(output, data_only=False)
            try:
                sheet = workbook["分类统计"]
                self.assertEqual(len(sheet._charts), 1)
                self.assertEqual(sheet["E22"].value, "本年度无可绘制数据")
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
