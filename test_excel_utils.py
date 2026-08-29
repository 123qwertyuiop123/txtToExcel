import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from excel_utils import save_workbook_safely, set_safe_text
from expense_converter import create_workbook, parse_money_items
from game_converter import create_game_workbook
from models import DayRecord, GameRecord


class ExcelSafetyTests(unittest.TestCase):
    def test_formula_prefixes_are_stored_as_text(self):
        workbook = Workbook()
        sheet = workbook.active
        for row, value in enumerate(("=2+2", "+2+2", "-2+2", "@SUM(A1)"), start=1):
            set_safe_text(sheet.cell(row, 1), value)
            self.assertEqual(sheet.cell(row, 1).data_type, "s")

    def test_safe_save_replaces_target_only_after_success(self):
        with TemporaryDirectory() as directory:
            target = Path(directory) / "result.xlsx"
            target.write_bytes(b"old workbook")
            workbook = Workbook()
            workbook.active["A1"] = "new workbook"
            save_workbook_safely(workbook, target)
            loaded = load_workbook(target, read_only=True)
            try:
                self.assertEqual(loaded.active["A1"].value, "new workbook")
            finally:
                loaded.close()
            self.assertEqual(list(target.parent.glob(".*.tmp.xlsx")), [])

    def test_failed_save_keeps_old_target_and_cleans_temporary_file(self):
        class FailingWorkbook:
            def save(self, path):
                Path(path).write_bytes(b"incomplete")
                raise RuntimeError("模拟保存失败")

        with TemporaryDirectory() as directory:
            target = Path(directory) / "result.xlsx"
            target.write_bytes(b"original")
            with self.assertRaisesRegex(RuntimeError, "模拟保存失败"):
                save_workbook_safely(FailingWorkbook(), target)
            self.assertEqual(target.read_bytes(), b"original")
            self.assertEqual(list(target.parent.glob(".*.tmp.xlsx")), [])

    def test_expense_detail_formula_is_not_executable(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "expense.xlsx"
            record = DayRecord(1, "=2+2:10", parse_money_items("=2+2:10"))
            create_workbook({1: {1: record}}, 2025, output)
            workbook = load_workbook(output, read_only=False, data_only=False)
            try:
                cell = workbook["1月"]["B1"]
                self.assertEqual(cell.value, "=2+2:10")
                self.assertEqual(cell.data_type, "s")
            finally:
                workbook.close()

    def test_game_name_formula_is_not_executable(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "game.xlsx"
            record = GameRecord(date(2025, 1, 1), "=2+2", Decimal("10"))
            create_game_workbook({2025: [record]}, output)
            workbook = load_workbook(output, read_only=False, data_only=False)
            try:
                detail = workbook["2025年"]["B2"]
                summary = workbook["2025年"]["H2"]
                self.assertEqual((detail.value, detail.data_type), ("=2+2", "s"))
                self.assertEqual((summary.value, summary.data_type), ("=2+2", "s"))
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
