"""代码审查发现问题的回归测试；所有写入均限定在临时目录。"""

import os
import unittest
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from category_charts import MONTHLY_LAYOUT, next_chart_row
from excel_utils import save_workbook_safely, set_safe_text, validate_output_path
from expense_converter import convert_many, create_workbook, infer_year_month, parse_txt
from game_converter import convert_game_txt_files, create_game_workbook
from money_utils import round_one_decimal
from models import DayRecord, MoneyItem
from year_merge import _copy_worksheet, _extract_month_totals, merge_yearly_workbooks, validate_yearly_workbooks


class ConversionBoundaryTests(unittest.TestCase):
    def test_cli_rejects_days_that_would_be_dropped(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "2025年2月消费.txt"
            output = Path(directory) / "result.xlsx"
            source.write_text("29日:晚饭:10", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "2025年2月消费.txt"):
                convert_many([source], output, None, None)
            self.assertFalse(output.exists())

    def test_leap_year_accepts_february_29(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "2024年2月消费.txt"
            source.write_text("29日:晚饭:10", encoding="utf-8")
            with patch("expense_converter.create_workbook") as create:
                convert_many([source], Path(directory) / "result.xlsx", None, None)
            self.assertEqual(create.call_args.args[0][2][29].expense, Decimal("10"))

    def test_parser_rejects_day_zero_with_line_number(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "input.txt"
            source.write_text("0日:晚饭:10", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "第 1 行"):
                parse_txt(source)

    def test_year_zero_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "年份必须"):
            infer_year_month(Path("0000年1月消费.txt"), None, None)

    def test_direct_workbook_creation_rejects_invalid_month_before_saving(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "result.xlsx"
            output.write_bytes(b"existing")
            with self.assertRaisesRegex(ValueError, "年月无效"):
                create_workbook({13: {}}, 2025, output)
            self.assertEqual(output.read_bytes(), b"existing")

    def test_wrong_output_extension_keeps_source_intact(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "2025年1月消费.txt"
            source.write_text("1日:晚饭:10", encoding="utf-8")
            original = source.read_bytes()
            with self.assertRaisesRegex(ValueError, "扩展名"):
                convert_many([source], source, None, None)
            self.assertEqual(source.read_bytes(), original)

    def test_game_converter_rejects_wrong_output_extension(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "2025年游戏.txt"
            source.write_text("1月1日:原神:10", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "扩展名"):
                convert_game_txt_files([source], source)
            self.assertEqual(source.read_text(encoding="utf-8"), "1月1日:原神:10")

    def test_empty_game_workbook_does_not_create_circular_total_formula(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "result.xlsx"
            for records in ({}, {2025: []}):
                with self.subTest(records=records), self.assertRaisesRegex(ValueError, "必须包含"):
                    create_game_workbook(records, output)
            self.assertFalse(output.exists())

    def test_excessively_large_amount_reports_a_user_error(self):
        with self.assertRaisesRegex(ValueError, "金额超出可处理范围"):
            round_one_decimal(Decimal("9" * 100))


class SavingBoundaryTests(unittest.TestCase):
    def test_output_cannot_replace_a_source_workbook(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "2025年消费统计.xlsx"
            source.write_bytes(b"original")
            with self.assertRaisesRegex(ValueError, "不能覆盖输入文件"):
                merge_yearly_workbooks([source], source)
            self.assertEqual(source.read_bytes(), b"original")

    def test_hard_link_alias_cannot_replace_source(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            alias = Path(directory) / "alias.xlsx"
            source.write_bytes(b"original")
            try:
                os.link(source, alias)
            except OSError as exc:
                self.skipTest(f"文件系统不支持硬链接：{exc}")
            with self.assertRaisesRegex(ValueError, "不能覆盖输入文件"):
                validate_output_path(alias, [source])

    def test_failed_replace_keeps_old_target_and_cleans_temp(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "result.xlsx"
            output.write_bytes(b"original")
            workbook = Workbook()
            with patch("excel_utils.os.replace", side_effect=PermissionError("文件被占用")):
                with self.assertRaisesRegex(PermissionError, "文件被占用"):
                    save_workbook_safely(workbook, output)
            self.assertEqual(output.read_bytes(), b"original")
            self.assertEqual(list(output.parent.glob(".*.tmp.xlsx")), [])


class YearMergeBoundaryTests(unittest.TestCase):
    def make_month(self, daily=10, income="=100+20-5", label="收入总计"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "1月"
        sheet["C1"] = daily
        sheet["B2"] = "总计"
        sheet["B3"] = label
        sheet["C3"] = income
        return workbook

    def test_literal_income_sum_is_supported(self):
        self.assertEqual(_extract_month_totals(self.make_month(), 1), (10, 115, 105))

    def test_reference_formula_is_not_misread_as_numbers(self):
        with self.assertRaisesRegex(ValueError, "C3.*无法可靠计算"):
            _extract_month_totals(self.make_month(income="=SUM(C1:C2)"), 1)

    def test_daily_formula_is_not_silently_omitted(self):
        with self.assertRaisesRegex(ValueError, "1月!C1"):
            _extract_month_totals(self.make_month(daily="=2+3"), 1)

    def test_boolean_and_nonfinite_amounts_are_rejected(self):
        for daily in (True, float("inf"), float("nan")):
            with self.subTest(daily=daily), self.assertRaises(ValueError):
                _extract_month_totals(self.make_month(daily=daily), 1)

    def test_digits_in_income_source_name_are_not_double_counted(self):
        book = self.make_month(income="=100", label="补贴+200:+100")
        self.assertEqual(_extract_month_totals(book, 1), (10, 100, 90))

    def test_copy_preserves_literal_formula_prefix_text(self):
        source = Workbook().active
        target = Workbook().active
        set_safe_text(source["A1"], "='1月'!C1")
        source["B1"] = "='1月'!C1"
        _copy_worksheet(source, target, 2025)
        self.assertEqual(target["A1"].value, "='1月'!C1")
        self.assertEqual(target["A1"].data_type, "s")
        self.assertEqual(target["B1"].value, "='2025-1月'!C1")
        self.assertEqual(target["B1"].data_type, "f")

    def test_multiple_years_merge_and_validation_catches_bad_formula(self):
        with TemporaryDirectory() as directory:
            sources = []
            for year in (2024, 2025):
                path = Path(directory) / f"{year}年消费统计.xlsx"
                records = {1: {1: DayRecord(1, "", [
                    MoneyItem("晚饭", Decimal("10"), False),
                    MoneyItem("工资", Decimal("100"), True),
                ])}}
                create_workbook(records, year, path)
                sources.append(path)
            output = Path(directory) / "年份统计.xlsx"
            self.assertEqual(validate_yearly_workbooks(sources), [])
            merge_yearly_workbooks(sources, output)
            result = load_workbook(output)
            try:
                self.assertEqual(result.sheetnames, ["年份统计", "2024年", "2025年"])
                self.assertEqual([result["年份统计"].cell(4, col).value for col in range(2, 5)], [20, 200, 180])
            finally:
                result.close()
            source = load_workbook(sources[0])
            try:
                source["1月"]["B33"] = "收入总计"
                source["1月"]["C33"] = "=SUM(C1:C31)"
                source["总计"]["A2"] = "13月"
                source.save(sources[0])
            finally:
                source.close()
            message = "\n".join(validate_yearly_workbooks(sources))
            self.assertIn(sources[0].name, message)
            self.assertIn("1月!C33", message)
            self.assertIn("月份无效：13月", message)


class ChartLayoutTests(unittest.TestCase):
    def test_dense_chart_reserves_enough_rows(self):
        sheet = Workbook().active
        self.assertEqual(next_chart_row(sheet, 2, 2, MONTHLY_LAYOUT), 18)
        self.assertGreater(next_chart_row(sheet, 2, 24, MONTHLY_LAYOUT), 20)


if __name__ == "__main__":
    unittest.main()
