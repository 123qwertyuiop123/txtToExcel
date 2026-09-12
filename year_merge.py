"""多个年度消费统计 Excel 的检查与合并。"""

import re
from collections import Counter
from copy import copy
from decimal import Decimal
from pathlib import Path
from typing import Iterable

from excel_utils import save_workbook_safely, set_safe_text, validate_output_path
from money_utils import decimal_to_number, round_one_decimal


YEAR_WORKBOOK_RE = re.compile(r"(?P<year>\d{4})\s*年.*统计")
LITERAL_SUM_RE = re.compile(r"[+-]?\d+(?:\.\d+)?(?:[+-]\d+(?:\.\d+)?)*")


# ---------- 输入工作簿检查 ----------

def infer_year_from_workbook_name(path: Path) -> int:
    """从年度统计工作簿文件名中提取年份。"""
    match = YEAR_WORKBOOK_RE.search(path.stem)
    if match is None:
        raise ValueError(f"无法从文件名识别年份：{path.name}\n建议命名为：2025年消费统计.xlsx")
    return int(match.group("year"))


def _summary_layout_errors(sheet, year: int) -> list[str]:
    """确认总表月份完整且唯一，避免复制出缺月份、重复月份或错误年份。"""
    labels = [str(row[0] or "") for row in sheet.iter_rows(min_col=1, max_col=1, values_only=True)]
    months = Counter(int(match.group(1)) for label in labels if (match := re.fullmatch(r"(\d{1,2})月", label)))
    errors = []
    for month in range(1, 13):
        if months[month] != 1:
            errors.append(f"“总计”工作表的 {month}月必须恰好出现一次，实际为 {months[month]} 次")
    for month in months:
        if not 1 <= month <= 12:
            errors.append(f"“总计”工作表月份无效：{month}月")
    annual_labels = [label for label in labels if label.endswith("年总计")]
    if annual_labels != [f"{year}年总计"]:
        errors.append(f"“总计”工作表必须恰好包含一个 {year}年总计 行，且年份与文件名一致")
    return errors


def validate_yearly_workbooks(paths: Iterable[Path]) -> list[str]:
    """检查年度工作簿的格式、年份唯一性和必要工作表。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ["缺少 openpyxl，请先执行：pip install -r requirements.txt"]

    errors: list[str] = []
    seen_years: dict[int, Path] = {}
    for path in paths:
        if path.suffix.lower() != ".xlsx":
            errors.append(f"文件格式不正确：{path.name}\n请选择 .xlsx 文件")
            continue
        if not path.is_file():
            errors.append(f"文件不存在：{path}")
            continue
        try:
            year = infer_year_from_workbook_name(path)
        except ValueError as exc:
            errors.append(str(exc))
            continue
        if year in seen_years:
            errors.append(f"年份重复：{year}年\n{seen_years[year].name}\n{path.name}")
            continue
        seen_years[year] = path
        # 损坏、加密或并非真正 xlsx 的文件可能抛出不同异常；在检查边界统一转为用户提示。
        try:
            workbook = load_workbook(path, read_only=True, data_only=False)
        except Exception as exc:
            errors.append(f"文件：{path.name}\n无法打开工作簿：{exc}")
            continue
        try:
            required = [*(f"{month}月" for month in range(1, 13)), "总计"]
            missing = [name for name in required if name not in workbook.sheetnames]
            if missing:
                errors.append(f"文件：{path.name}\n缺少工作表：{'、'.join(missing)}")
            if not missing:
                errors.extend(f"文件：{path.name}\n{error}" for error in _summary_layout_errors(workbook["总计"], year))
                # 检查阶段也确认金额可以可靠提取，避免“检查通过”后才发现复杂公式。
                for month in range(1, 13):
                    try:
                        _extract_month_totals(workbook, month)
                    except ValueError as exc:
                        errors.append(f"文件：{path.name}\n{exc}")
        finally:
            workbook.close()
    if not seen_years and not errors:
        errors.append("没有选择年度消费统计 Excel 文件")
    return errors


# ---------- 工作表复制与数值提取 ----------

def _rewrite_year_formula(value: object, year: int) -> object:
    """为复制后的跨表公式补充年份前缀。"""
    if not isinstance(value, str) or not value.startswith("="):
        return value
    return re.sub(
        r"'?(?P<month>\d{1,2}月)'?!",
        lambda match: f"'{year}-{match.group('month')}'!",
        value,
    )


def _copy_worksheet(source, target, year: int) -> None:
    """复制工作表的值、样式、尺寸和合并单元格。"""
    for row in source.iter_rows():
        for source_cell in row:
            target_cell = target[source_cell.coordinate]
            # 值以“=”开头不等于公式；复制时必须保留来源单元格的文本类型。
            if source_cell.data_type == "f":
                target_cell.value = _rewrite_year_formula(source_cell.value, year)
            elif isinstance(source_cell.value, str):
                set_safe_text(target_cell, source_cell.value)
            else:
                target_cell.value = source_cell.value
            if source_cell.has_style:
                target_cell.font = copy(source_cell.font)
                target_cell.fill = copy(source_cell.fill)
                target_cell.border = copy(source_cell.border)
                target_cell.alignment = copy(source_cell.alignment)
                target_cell.number_format = source_cell.number_format
                target_cell.protection = copy(source_cell.protection)
            if source_cell.hyperlink:
                target_cell._hyperlink = copy(source_cell.hyperlink)
            if source_cell.comment:
                target_cell.comment = copy(source_cell.comment)
    for key, source_dimension in source.column_dimensions.items():
        target_dimension = target.column_dimensions[key]
        target_dimension.width = source_dimension.width
        target_dimension.hidden = source_dimension.hidden
        target_dimension.bestFit = source_dimension.bestFit
    for index, source_dimension in source.row_dimensions.items():
        target_dimension = target.row_dimensions[index]
        target_dimension.height = source_dimension.height
        target_dimension.hidden = source_dimension.hidden
    for merged_range in source.merged_cells.ranges:
        target.merge_cells(str(merged_range))
    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.sheet_format.defaultColWidth = source.sheet_format.defaultColWidth
    target.sheet_format.defaultRowHeight = source.sheet_format.defaultRowHeight


def _numeric_amount(value: object, location: str) -> Decimal:
    """只接受有限数值；布尔值、错误值和待计算公式不能当作金额。"""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{location}不是可读取的金额：{value}")
    amount = Decimal(str(value))
    if not amount.is_finite():
        raise ValueError(f"{location}金额不是有限数值")
    return amount


def _extract_month_totals(workbook, month: int) -> tuple[int | float, int | float, int | float]:
    """从月度明细重新计算支出和收入，避免依赖 Excel 公式缓存。"""
    sheet = workbook[f"{month}月"]
    expense = Decimal("0")
    # 流式扫描 B:C；read_only 模式下逐个调用 cell 会反复读取文件，成本很高。
    rows = enumerate(sheet.iter_rows(min_col=2, max_col=3, values_only=True), start=1)
    for row, (label, value) in rows:
        if label in ("总计", "支出总计"):
            total_row = row
            break
        if value is not None:
            expense += _numeric_amount(value, f"{month}月!C{row}")
    else:
        raise ValueError(f"工作表“{month}月”中找不到总计行")
    income_row = total_row + 1
    _, (income_label, income_formula) = next(rows, (income_row, (None, None)))
    income_text = str(income_label or "")
    # 只从“项目:+金额”提取收入，名称中的“+数字”不是一笔收入。
    income_values = [Decimal(value) for value in re.findall(r"[:：]\s*\+\s*(\d+(?:\.\d+)?)", income_text)]
    # 优先读取收入说明文字中的“+金额”；兼容旧工作簿时再尝试数值或公式。
    if income_values:
        income = sum(income_values, Decimal("0"))
    else:
        if isinstance(income_formula, (int, float)):
            income = _numeric_amount(income_formula, f"{month}月!C{income_row}")
        elif isinstance(income_formula, str) and income_formula.startswith("="):
            expression = income_formula[1:].replace(" ", "")
            # 仅支持本程序写出的字面量加减式；不能从 SUM(C1:C3) 中提取 1 和 3。
            # 不调用 eval，也不声称能够计算任意 Excel 公式。
            if not LITERAL_SUM_RE.fullmatch(expression):
                raise ValueError(f"{month}月!C{income_row}收入公式无法可靠计算：{income_formula}；请改为金额或补充收入明细")
            numbers = re.findall(r"[+-]?\d+(?:\.\d+)?", expression)
            income = sum((Decimal(value) for value in numbers), Decimal("0"))
        elif income_formula is None:
            income = Decimal("0")
        else:
            raise ValueError(f"{month}月!C{income_row}收入金额无法识别：{income_formula}")

    expense = round_one_decimal(expense)
    income = round_one_decimal(income)
    balance = round_one_decimal(income - expense)
    return decimal_to_number(expense), decimal_to_number(income), decimal_to_number(balance)


# ---------- 合并输出 ----------

def merge_yearly_workbooks(paths: Iterable[Path], output_path: Path) -> Path:
    """只复制各年度总情况，并生成跨年份汇总。"""
    source_paths = list(paths)
    output_path = validate_output_path(output_path, source_paths)
    errors = validate_yearly_workbooks(source_paths)
    if errors:
        raise ValueError("\n\n".join(errors))
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请先执行：pip install -r requirements.txt") from exc

    sources = sorted((infer_year_from_workbook_name(path), path) for path in source_paths)
    combined = Workbook()
    summary = combined.active
    summary.title = "年份统计"
    summary.append(["年份", "支出", "收入", "结余（收入-支出）"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    yearly_totals: list[tuple[int, int | float, int | float, int | float]] = []
    # 每次只打开一个来源工作簿，完成提取后立即关闭，降低多年度合并时的资源占用。
    for year, path in sources:
        source_book = load_workbook(path, data_only=False)
        try:
            try:
                month_totals = {month: _extract_month_totals(source_book, month) for month in range(1, 13)}
            except ValueError as exc:
                raise ValueError(f"文件：{path.name}\n{exc}") from exc
            annual_expense = decimal_to_number(round_one_decimal(sum(
                (Decimal(str(values[0])) for values in month_totals.values()), Decimal("0")
            )))
            annual_income = decimal_to_number(round_one_decimal(sum(
                (Decimal(str(values[1])) for values in month_totals.values()), Decimal("0")
            )))
            annual_balance = decimal_to_number(round_one_decimal(
                Decimal(str(annual_income)) - Decimal(str(annual_expense))
            ))

            # 复制原“总计”样式，但把跨表公式替换成刚刚重新计算的静态数值。
            target_sheet = combined.create_sheet(f"{year}年")
            _copy_worksheet(source_book["总计"], target_sheet, year)
            total_row: int | None = None
            for row in range(1, target_sheet.max_row + 1):
                label = target_sheet.cell(row, 1).value
                month_match = re.fullmatch(r"(\d{1,2})月", str(label or ""))
                if month_match:
                    for column, value in enumerate(month_totals[int(month_match.group(1))], start=2):
                        target_sheet.cell(row, column, value)
                elif isinstance(label, str) and label.endswith("年总计"):
                    total_row = row
                    target_sheet.cell(row, 1, f"{year}年总计")
                    target_sheet.cell(row, 2, annual_expense)
                    target_sheet.cell(row, 3, annual_income)
                    target_sheet.cell(row, 4, annual_balance)
            if total_row is None:
                raise ValueError(f"文件：{path.name}\n找不到年度总计行")
            yearly_totals.append((year, annual_expense, annual_income, annual_balance))
        finally:
            source_book.close()

    for row, (year, expense, income, balance) in enumerate(yearly_totals, start=2):
        summary.cell(row, 1, f"{year}年")
        summary.cell(row, 2, expense)
        summary.cell(row, 3, income)
        summary.cell(row, 4, balance)
    final_row = len(yearly_totals) + 2
    summary.cell(final_row, 1, "全部年份")
    total_expense = round_one_decimal(sum(
        (Decimal(str(values[1])) for values in yearly_totals), Decimal("0")
    ))
    total_income = round_one_decimal(sum(
        (Decimal(str(values[2])) for values in yearly_totals), Decimal("0")
    ))
    summary.cell(final_row, 2, decimal_to_number(total_expense))
    summary.cell(final_row, 3, decimal_to_number(total_income))
    summary.cell(final_row, 4, decimal_to_number(round_one_decimal(total_income - total_expense)))
    for cell in summary[final_row]:
        cell.font = Font(bold=True)
    for column, width in {"A": 14, "B": 18, "C": 18, "D": 24}.items():
        summary.column_dimensions[column].width = width
    summary.freeze_panes = "A2"
    combined.calculation.fullCalcOnLoad = True
    combined.calculation.forceFullCalc = True
    save_workbook_safely(combined, output_path)
    return output_path
